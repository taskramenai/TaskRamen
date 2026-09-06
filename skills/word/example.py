#!/usr/bin/env python3
# TaskRamen.ai
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Jobs Jolt Private Limited (TaskRamen.ai)
"""Starter template for building a .docx with python-docx (MIT-licensed).

Copy into the active project folder and adapt. Run:
    pip install --quiet python-docx
    python3 example.py
Produces doc.docx in the same folder.
"""
import sys
from docx import Document
from docx.shared import Pt

OUT = sys.argv[1] if len(sys.argv) > 1 else "doc.docx"

doc = Document()                          # or Document("template.docx") for branded output
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)

doc.add_heading("Service Proposal", level=0)

doc.add_heading("Overview", level=1)
doc.add_paragraph(
    "Acme will deliver a marketing analytics package covering ad performance, "
    "lead tracking, and a weekly business scorecard."
)

doc.add_heading("Scope", level=1)
for item in ["Google Ads + Meta Ads reporting", "Lead-list research", "Monthly review call"]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Pricing", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text = "Item", "Monthly (USD)"
for name, price in [("Ads management", "1,200"), ("Reporting", "400")]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = name, price

doc.save(OUT)
print(f"Wrote {OUT}")
