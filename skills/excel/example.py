#!/usr/bin/env python3
# TaskRamen.ai
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Jobs Jolt Private Limited (TaskRamen.ai)
"""Starter template for building a .xlsx workbook with openpyxl (MIT-licensed).

Copy into the active project folder and adapt. Run:
    pip install --quiet openpyxl
    python3 example.py
Produces book.xlsx in the same folder. Uses live formulas, not hardcoded results.
"""
import sys
import openpyxl
from openpyxl.styles import Font, Alignment

OUT = sys.argv[1] if len(sys.argv) > 1 else "book.xlsx"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "P&L"

headers = ["Month", "Revenue", "Costs", "Profit"]
ws.append(headers)
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

rows = [("Jan", 12000, 8000), ("Feb", 14500, 8200), ("Mar", 15300, 8600)]
for r, (month, rev, cost) in enumerate(rows, start=2):
    ws.cell(r, 1, month)
    ws.cell(r, 2, rev)
    ws.cell(r, 3, cost)
    ws.cell(r, 4, f"=B{r}-C{r}")          # live formula

total_row = len(rows) + 2
ws.cell(total_row, 1, "Total").font = Font(bold=True)
for col in ("B", "C", "D"):
    ws[f"{col}{total_row}"] = f"=SUM({col}2:{col}{total_row-1})"
    ws[f"{col}{total_row}"].font = Font(bold=True)

# Currency format, negatives in parentheses
for col in ("B", "C", "D"):
    for cell in ws[col][1:]:
        cell.number_format = '#,##0;(#,##0)'

ws.freeze_panes = "A2"
for col, width in {"A": 10, "B": 14, "C": 14, "D": 14}.items():
    ws.column_dimensions[col].width = width

wb.save(OUT)
print(f"Wrote {OUT} (open once in Excel/LibreOffice to populate formula results)")
