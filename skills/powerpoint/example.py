#!/usr/bin/env python3
# TaskRamen.ai
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Jobs Jolt Private Limited (TaskRamen.ai)
"""Starter template for building a .pptx deck with python-pptx (MIT-licensed).

Copy into the active project folder and adapt. Run:
    pip install --quiet python-pptx
    python3 example.py
Produces deck.pptx in the same folder.
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

OUT = sys.argv[1] if len(sys.argv) > 1 else "deck.pptx"
BRAND = RGBColor(0x1F, 0x49, 0xE0)  # swap for the user's brand color

prs = Presentation()                 # or Presentation("template.pptx") for a branded deck
prs.slide_width = Inches(13.333)     # 16:9
prs.slide_height = Inches(7.5)

# 1) Title slide
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Quarterly Business Review"
s.placeholders[1].text = "Acme Pte Ltd  -  Q2 2026"

# 2) Bulleted highlights slide
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Highlights"
tf = s.placeholders[1].text_frame
tf.text = "Revenue up 18% QoQ"
for line in ["3 new enterprise logos", "Churn down to 1.2%", "NPS at 61"]:
    p = tf.add_paragraph()
    p.text = line

# 3) Native chart slide (prefer this over pasted screenshots)
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Revenue by Quarter"
data = CategoryChartData()
data.categories = ["Q1", "Q2", "Q3", "Q4"]
data.add_series("Revenue ($k)", (210, 248, 0, 0))
s.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(1), Inches(1.8), Inches(11), Inches(5), data,
)

prs.save(OUT)
print(f"Wrote {OUT}")
